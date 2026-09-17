"""
Builds a deliberately nasty, large workbook to prove the tracer finds real
problems in something resembling a 100k-variable model.

Planted faults:
  1. out_capex depends on a hard-coded cell nobody told us about
  2. out_gypsum_pct depends on NO input at all (a frozen constant)
  3. out_unit_cost references an EMPTY cell (broken string)
  4. out_npv chain goes through INDIRECT (untraceable)
  5. in_credit_multiplier is wired but drives nothing
  6. Two enormous sheets of irrelevant working (the 99% we want to discard)
"""
import openpyxl
from openpyxl.utils import absolute_coordinate, quote_sheetname
from openpyxl.workbook.defined_name import DefinedName

wb = openpyxl.Workbook()

def name(n, sheet, cell):
    wb.defined_names.add(DefinedName(
        n, attr_text=f"{quote_sheetname(sheet)}!{absolute_coordinate(cell)}"))

# ---------------- Inputs ----------------
ws = wb.active; ws.title = "Inputs"
rows = [("Feedstock carbon",1.0,"in_feedstock_carbon"),("Feedstock yield",1.0,"in_feedstock_yield"),
        ("Feedstock cost",180,"in_feedstock_cost"),("Plant capacity",5000000,"in_plant_capacity"),
        ("Gas consumption",9.5,"in_gas_consumption"),("Line conversion",40,"in_line_conversion"),
        ("Biochar rate",15,"in_biochar_rate"),("Carbon price",120,"in_carbon_price"),
        ("Credit multiplier",1.0,"in_credit_multiplier")]
for i,(lbl,val,nm) in enumerate(rows, start=2):
    ws[f"A{i}"]=lbl; ws[f"B{i}"]=val; name(nm,"Inputs",f"B{i}")

# ---------------- Hidden assumptions sheet (the trap) ----------------
wsa = wb.create_sheet("Assumptions")
wsa["A1"]="Regional grid factor"; wsa["B1"]=0.233      # hard-coded, not an input
wsa["A2"]="Plant derate"; wsa["B2"]=0.87               # hard-coded, not an input
wsa["A3"]="Legacy gypsum pct"; wsa["B3"]=21.75         # frozen constant

# ---------------- Calculations ----------------
wsc = wb.create_sheet("Calcs")
wsc["B1"]="=in_line_conversion/100"
wsc["B2"]="=in_biochar_rate/100"
wsc["B3"]="=in_plant_capacity*B1"
wsc["B4"]="=in_plant_capacity*B1*B2*0.00065*(1/in_feedstock_yield)"
wsc["B5"]="=B4*in_feedstock_cost"
wsc["B6"]="=165000*(in_plant_capacity/5000000)"
wsc["B7"]="=(in_plant_capacity*in_gas_consumption*B1*0.45/1000)*38"
wsc["B8"]="=(in_plant_capacity*B1*B2*0.0091*in_feedstock_carbon)*in_carbon_price"
wsc["B9"]="=INDIRECT(\"Calcs!B8\")*1.0"   # FAULT 4: untraceable
wsc["B10"]="=B5+B6"

# ---------------- Two huge irrelevant sheets (the 99%) ----------------
for sheet_name in ("LegacyModel2019","SensitivityDump"):
    wsj = wb.create_sheet(sheet_name)
    for r in range(1, 601):
        for c in range(1, 21):
            wsj.cell(row=r, column=c, value=(f"=A{max(1,r-1)}*1.01" if c==1 and r>1 else r*c*0.37))

# ---------------- Outputs ----------------
wso = wb.create_sheet("Outputs")
outs = [
 ("out_gas_pct",     "=MIN(100,Calcs!B1*62)"),
 ("out_gas_mwh",     "=in_plant_capacity*in_gas_consumption*Calcs!B1*0.45/1000"),
 ("out_gypsum_pct",  "=Assumptions!B3"),                                   # FAULT 2
 ("out_gypsum_tonnes","=in_plant_capacity*Calcs!B1*Calcs!B2*0.052*in_feedstock_yield"),
 ("out_net_carbon",  "=6.4-(Calcs!B1*Calcs!B2*22*in_feedstock_carbon)"),
 ("out_cdr_tonnes",  "=in_plant_capacity*Calcs!B1*Calcs!B2*0.0091*in_feedstock_carbon"),
 ("out_capex",       "=2600000*(in_plant_capacity/5000000)*(0.6+Calcs!B1*0.4)*Assumptions!B2"), # FAULT 1
 ("out_opex",        "=Calcs!B10"),
 ("out_unit_cost",   "=IF(Calcs!B3>0,Calcs!B10/Calcs!B3,0)+Calcs!B47"),    # FAULT 3 empty cell
 ("out_net_cash",    "=Calcs!B7+Calcs!B8-Calcs!B10"),
 ("out_payback",     "=IF(out_net_cash>0,out_capex/out_net_cash,-1)"),
 ("out_npv",         "=-out_capex+Calcs!B9*((1-(1.08^-10))/0.08)"),        # via INDIRECT
]
for i,(nm,f) in enumerate(outs, start=2):
    wso[f"A{i}"]=nm; wso[f"B{i}"]=f; name(nm,"Outputs",f"B{i}")

wb.save("messy_model.xlsx")
tot = sum(wb[s].max_row*wb[s].max_column for s in wb.sheetnames)
print(f"messy_model.xlsx built, approx {tot:,} cells across {len(wb.sheetnames)} sheets")
