"""The startup self-test, the safety net for Adaptavate changing their model."""

import base64
import shutil

import openpyxl
import pytest

from app.calc import SelfTestFailed, build_engine


def test_engine_builds_on_a_good_workbook(test_model):
    engine = build_engine(test_model, backend="formulas")
    assert engine.backend_name == "formulas"


def test_missing_workbook_raises(tmp_path):
    with pytest.raises(RuntimeError, match="not found"):
        build_engine(tmp_path / "nope.xlsx", backend="formulas")


def test_base64_encoded_workbook_still_builds(test_model, tmp_path):
    """Some hosting dashboards' 'secret file' box is a text paste field, which
    mangles a raw .xlsx. The documented workaround is to paste base64 text
    instead; the engine must transparently decode it. See app/calc/engine.py,
    _materialize_workbook.
    """
    encoded = tmp_path / "workbook.b64"
    encoded.write_bytes(base64.b64encode(test_model.read_bytes()))

    engine = build_engine(encoded, backend="formulas")
    assert engine.backend_name == "formulas"


def test_corrupted_workbook_raises_a_clear_error(tmp_path):
    """Neither a real .xlsx nor valid base64: the classic symptom of a binary
    file pasted into a text-only field. Must not be a confusing traceback.
    """
    garbage = tmp_path / "corrupted.xlsx"
    garbage.write_bytes(b"not a zip file and not base64 either !!! \xff\xfe")

    with pytest.raises(RuntimeError, match="neither a valid .xlsx nor valid base64"):
        build_engine(garbage, backend="formulas")


def test_changed_model_refuses_to_start(test_model, tmp_path):
    """Alter a hidden coefficient: the service must refuse to boot."""
    changed = tmp_path / "changed.xlsx"
    shutil.copy(test_model, changed)
    wb = openpyxl.load_workbook(changed)
    wb["Calculations"]["B8"] = "=195000*(in_plant_capacity/5000000)"
    wb.save(changed)
    wb.close()

    with pytest.raises(SelfTestFailed, match="self-test failed"):
        build_engine(changed, backend="formulas")


def test_missing_named_range_refuses_to_start(test_model, tmp_path):
    from app.calc.backends.formulas_backend import MissingNamedRange

    broken = tmp_path / "broken.xlsx"
    shutil.copy(test_model, broken)
    wb = openpyxl.load_workbook(broken)
    del wb.defined_names["out_npv"]
    wb.save(broken)
    wb.close()

    with pytest.raises(MissingNamedRange, match="out_npv"):
        build_engine(broken, backend="formulas")


def test_inserted_row_does_not_break_the_wiring(test_model, tmp_path):
    """Named ranges follow the cell. This is why we refuse coordinates.

    openpyxl does not shift defined names on insert_rows the way Excel does, so
    the shift is applied explicitly to reproduce real Excel behaviour.
    """
    from openpyxl.workbook.defined_name import DefinedName

    edited = tmp_path / "edited.xlsx"
    shutil.copy(test_model, edited)
    wb = openpyxl.load_workbook(edited)
    wb["Inputs"].insert_rows(2)
    wb["Inputs"]["A2"] = "Model version"
    wb["Inputs"]["B2"] = "v2.4"
    for name in list(wb.defined_names):
        ref = wb.defined_names[name].attr_text
        if "'Inputs'!" in ref:
            col = ref.split("!")[1].replace("$", "")[0]
            row = int(ref.split("$")[-1])
            del wb.defined_names[name]
            wb.defined_names.add(DefinedName(name, attr_text=f"'Inputs'!${col}${row + 1}"))
    wb.save(edited)
    wb.close()

    # Must still build, and still produce the same answers.
    engine = build_engine(edited, backend="formulas")
    from app.calc.contract import SELF_TEST_EXPECTED, SELF_TEST_INPUT
    result = engine.calculate(SELF_TEST_INPUT)
    for key, expected in SELF_TEST_EXPECTED.items():
        assert result[key] == pytest.approx(expected, rel=1e-4)
